"""把上传文件统一解析为 Markdown。页码用 <!-- p:N --> 标记，供分块时写入原文定位。"""
import csv
import io
from html.parser import HTMLParser

SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".csv", ".html", ".htm"}


class ParseError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文本编码")


def _md_table(rows: list[list[str]]) -> str:
    rows = [[(c or "").replace("|", "\\|").replace("\n", " ").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "-|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        parts.append(f"<!-- p:{i} -->\n{(page.extract_text() or '').strip()}")
    text = "\n\n".join(parts)
    if len(text.replace("<!-- p:", "")) < 20:
        raise ParseError("PDF 未提取到文字，可能是扫描件；扫描件需 OCR，演示版未开放")
    return text


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    out = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph

            p = Paragraph(block, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") and style[-1:].isdigit():
                out.append("#" * int(style[-1]) + " " + text)
            elif style == "title":
                out.append("# " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            from docx.table import Table

            t = Table(block, doc)
            out.append(_md_table([[c.text for c in row.cells] for row in t.rows]))
    return "\n\n".join(out)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        table = _md_table(rows)
        if table:
            out.append(f"## {ws.title}\n\n{table}")
    return "\n\n".join(out)


def _csv(data: bytes) -> str:
    return _md_table(list(csv.reader(io.StringIO(_decode(data)))))


class _Html(HTMLParser):
    SKIP = {"script", "style", "nav", "footer", "header", "aside"}

    def __init__(self):
        super().__init__()
        self.out: list[str] = []
        self.skip = 0
        self.prefix = ""

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in {"h1", "h2", "h3", "h4"}:
            self.out.append("\n\n")
            self.prefix = "#" * int(tag[1]) + " "
        elif tag in {"p", "div", "li", "br", "tr"}:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in {"h1", "h2", "h3", "h4"}:
            self.out.append("\n\n")

    def handle_data(self, data):
        if self.skip or not data.strip():
            return
        self.out.append(self.prefix + data.strip())
        self.prefix = ""


def _html(data: bytes) -> str:
    h = _Html()
    h.feed(_decode(data))
    lines = [ln.strip() for ln in "".join(h.out).splitlines()]
    return "\n".join(ln for ln in lines if ln)


def read_rows(filename: str, data: bytes) -> list[tuple[str, list[list[str]]]]:
    """表格文件的原始行，供结构化抽取使用：[(sheet 名, 行列表)]。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "csv":
        return [("CSV", [[c.strip() for c in r] for r in csv.reader(io.StringIO(_decode(data))) if any(r)])]
    if ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [
                ["" if v is None else str(v).strip() for v in row]
                for row in ws.iter_rows(max_row=5000, max_col=50, values_only=True)
            ]
            rows = [r for r in rows if any(r)]
            if rows:
                out.append((ws.title, rows))
        return out
    raise ParseError("数据表格只支持 xlsx 或 csv 文件")


def parse(filename: str, data: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED:
        raise ParseError(f"暂不支持 {ext or '无扩展名'} 文件；支持：{'、'.join(sorted(SUPPORTED))}")
    if ext in {".md", ".txt"}:
        md = _decode(data)
    elif ext == ".pdf":
        md = _pdf(data)
    elif ext == ".docx":
        md = _docx(data)
    elif ext == ".xlsx":
        md = _xlsx(data)
    elif ext == ".csv":
        md = _csv(data)
    else:
        md = _html(data)
    md = md.strip()
    if not md:
        raise ParseError("文件中没有可提取的文字")
    return md
